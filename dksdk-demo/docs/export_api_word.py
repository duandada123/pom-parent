"""Generate the local API reference from swagger.json using python-docx."""
import json
from pathlib import Path
from urllib.request import urlopen
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
spec = json.loads((ROOT / 'swagger.json').read_text(encoding='utf-8-sig'))
try:
    with urlopen('http://localhost:8080/v2/api-docs', timeout=5) as response:
        live = json.load(response)
    assert live['paths'] == spec['paths'] and live['definitions'] == spec['definitions'], 'Live API differs from saved Swagger'
    print('Live Swagger matches saved API contract')
except OSError:
    print('Using saved Swagger contract')
models = spec['definitions']
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.62)
sec.left_margin = sec.right_margin = Inches(.65)
sec.footer_distance = Inches(.3)
for name, size in [('Normal', 10.5), ('Title', 24), ('Heading 1', 17), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = Pt(15 if name == 'Normal' else size * 1.25)
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); s.element.get_or_add_pPr().append(snap)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('DataWorks API  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for r in footer.runs: r.font.size = Pt(8)
doc.core_properties.title = 'DataWorks SDK 接口文档'
doc.core_properties.subject = '脚本创建 提交 发布和运行日志查询'
doc.core_properties.author = '项目开发文档'

def p(text): return doc.add_paragraph(text)
def h(text, level=2): return doc.add_heading(text, level)
def page(title):
    heading = h(title, 1)
    heading.paragraph_format.page_break_before = True
def code(value):
    text = json.dumps(value, ensure_ascii=False, indent=2) if not isinstance(value, str) else value
    if not isinstance(value, str):
        compact = json.dumps(value, ensure_ascii=False, separators=(',', ':'))
        if len(compact) <= 108:
            text = compact
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = Pt(11)
    para.paragraph_format.space_after = Pt(7)
    for i, line in enumerate(text.splitlines()):
        r = para.add_run(('\n' if i else '') + line)
        r.font.name = 'Consolas'; r.font.size = Pt(9)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    return para

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate([headers] + rows):
        row = t.rows[0] if i == 0 else t.add_row()
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '70'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            pr.append(margins)
            if i == 0 or i % 2 == 0:
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'E4EBF2' if i == 0 else 'F7F9FB'); pr.append(fill)
            para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = Pt(12.5)
            if j == 1 and len(headers) > 2: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(value)); run.font.size = Pt(9.5); run.bold = i == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(5)
    spacer.add_run().font.size = Pt(1)

def fields(model, request=False, names=None):
    definition = models[model]; props = definition['properties']; rows = []
    for name in names or props.keys():
        item = props[name]
        typ = item.get('type', 'object')
        typ = {'integer': 'Long' if item.get('format') == 'int64' else 'Integer', 'string':'String', 'boolean':'Boolean', 'array':'Array', 'object':'Object'}[typ]
        desc = item.get('description', '')
        if item.get('enum'): desc += '；可选值：' + '、'.join(item['enum'])
        if name == 'fileName' and model == 'CreateScriptRequest': desc += '；最长 255 字符'
        if request and name in ['projectId','fileId','deploymentId','taskId','instanceId','fileType','resourceGroupId']: desc += '；传入时必须大于 0'
        required = '必填' if name in definition.get('required', []) else '可选'
        if name == 'projectId': required = '条件'
        if model == 'DeployScriptRequest' and name == 'fileId': required = '条件'
        if model == 'QueryLogsRequest' and name in ['taskId','fileName','instanceId','bizdate']: required = '条件'
        if request: rows.append([name, typ, required, desc])
        else: rows.append([name, typ, desc])
    if request:
        table(['字段', '类型', '必填', '说明'], rows, [1.95, .65, .5, 4.1])
    else:
        table(['字段', '类型', '说明'], rows, [2.0, .85, 4.35])

def ok(data):
    return {'success':True,'code':'OK','message':'请求处理成功，业务状态请查看 data','requestId':None,'data':data}

doc.add_heading('DataWorks SDK 接口文档', 0)
p('dksdk-demo  |  接口版本 1.0.0  |  2026年10月22日')
p('本文供接口调用方联调使用，覆盖脚本创建、提交、发布和运行日志查询。提交与发布是异步操作，HTTP 200 或 success=true 仅表示本次请求处理正常，最终结果必须检查业务状态。')
h('1 接口总览', 1)
table(['用途', '方法', '接口路径'], [
    ['创建脚本','POST','/api/dataworks/files/create'],['提交脚本','POST','/api/dataworks/files/submit'],
    ['检查状态并发布','POST','/api/dataworks/files/deploy'],['查询运行日志','POST','/api/dataworks/logs/query']], [1.35,.7,5.15])
p('本地服务地址：http://localhost:8080。请求头 Content-Type: application/json；请求体和响应均使用 JSON。实际部署时替换主机及端口。')
p('在线文档：http://localhost:8080/swagger-ui.html\nSwagger 定义：http://localhost:8080/v2/api-docs')
h('推荐调用顺序')
p('创建脚本取得 fileId → 提交取得 deploymentId → 检查提交状态并发布 → 保存 publishDeploymentId → 使用 checkOnly=true 查询最终发布结果。发布不会自动运行任务；产生运行实例后才能查询日志。')
h('统一响应结构')
fields('ApiResponse«CreateScriptResponse»', names=['success','code','message','requestId','data'])
p('成功结果中的云端 requestId 通常位于 data；顶层 requestId 在成功时为 null。以下示例中的所有 ID 和日志内容均为示意值，调用时必须替换。')

page('2 创建脚本请求参数')
code('POST /api/dataworks/files/create')
p('仅创建文件并返回 fileId，不自动提交、发布或运行。projectId 可省略，但必须已配置有效的默认工作空间 ID。必填的字符串字段不能只包含空白。')
fields('CreateScriptRequest', True, ['projectId','fileName','fileType','fileFolderPath','content','fileDescription','connectionName','inputList','outputList','paraValue','cronExpress','cycleType','owner','resourceGroupIdentifier','resourceGroupId','createFolderIfNotExists','autoParsing','autoRerunTimes','autoRerunIntervalMillis'])
p('fileType 应按目标工作空间支持的类型选择；1221 仅沿用项目 Python 示例。输入输出依赖、数据源和调度资源组需按任务实际配置。')

page('3 创建示例与提交接口')
h('创建请求示例')
code({'projectId':10000,'fileName':'python_demo','fileType':1221,'fileFolderPath':'业务流程/demo/MaxCompute','content':"print('hello DataWorks')"})
h('创建响应数据')
fields('CreateScriptResponse')
p('成功响应的 data 示例：')
code({'fileId':1000001,'requestId':'aliyun-request-id'})
h('提交脚本')
code('POST /api/dataworks/files/submit')
p('提交返回提交包 deploymentId。state=SUBMIT_ACCEPTED 表示请求已受理，不表示已提交完成；后续使用发布接口检查。')
fields('SubmitScriptRequest', True, ['projectId','fileId','comment'])
code({'projectId':10000,'fileId':1000001,'comment':'提交示例'})
h('提交响应数据')
fields('SubmitScriptResponse', names=['fileId','deploymentId','state','requestId'])
code({'fileId':1000001,'deploymentId':3000001,'state':'SUBMIT_ACCEPTED','requestId':'aliyun-request-id'})

page('4 发布条件与状态处理')
code('POST /api/dataworks/files/deploy')
fields('DeployScriptRequest', True, ['projectId','fileId','deploymentId','checkOnly','comment'])
h('实际发布的前置条件')
p('checkOnly=false 时，deploymentId 必须是该文件的开发环境提交包。服务检查 toEnvironment=1、status=1、checkingStatus=null，并确认包内文件明细包含 fileId，全部满足才调用 DeployFile。无需额外传入 nodeId。')
h('业务状态与调用方处理')
table(['state', '含义及后续操作'], [
    ['WAITING','提交或检查未完成。稍后再检查；尚未发起发布。'],
    ['SUBMIT_FAILED','提交包 status=2。排查失败原因，不继续发布。'],
    ['CHECK_FAILED','发布前检查未通过，checkingStatus=7。修正检查问题。'],
    ['PUBLISH_ACCEPTED','发布已受理。保存 publishDeploymentId，后续仅查询新包。'],
    ['STATUS_ONLY','只读查询结果。结合 status、checkingStatus 和 toEnvironment 判断。']], [2.0,5.2])
p('status：0 就绪，1 成功，2 失败。checkingStatus：7 未通过，8 检查中；实际发布必须为空。toEnvironment：1 开发环境，2 生产环境。')
h('避免重复发布')
p('收到 PUBLISH_ACCEPTED 后，不要继续用原提交包反复请求发布。将新的 publishDeploymentId 作为 deploymentId，并设置 checkOnly=true 查询。服务未实现持久化幂等记录，调用方须保存关联 ID，避免并发重复发布或使用旧提交包发布已修改文件。')

page('5 发布响应与查询示例')
fields('DeployScriptResponse', names=['fileId','deploymentId','state','status','checkingStatus','toEnvironment','publishDeploymentId','requestId','message'])
h('发布请求示例')
code({'projectId':10000,'fileId':1000001,'deploymentId':3000001,'comment':'发布示例'})
p('发布受理后，data 示例：')
code({'fileId':1000001,'deploymentId':3000001,'state':'PUBLISH_ACCEPTED','status':1,'checkingStatus':None,'toEnvironment':1,'publishDeploymentId':4000001,'requestId':'aliyun-request-id','message':'发布请求已受理，请查询新发布包的最终状态'})
h('查询最终发布结果')
code({'projectId':10000,'deploymentId':4000001,'checkOnly':True})
p('返回 state=STATUS_ONLY。此时 status 对应 4000001 发布包；应结合 toEnvironment=2 确认生产发布结果。只读查询不会再次执行发布，也可以用来查看提交包状态。')

page('6 日志查询请求参数')
code('POST /api/dataworks/logs/query')
p('taskId、fileName、instanceId 必须且只能选择一个。已知 instanceId 时直接查日志；已知 taskId 时无需按名称查找。按任务或文件名查询必须提供 bizdate，直接按实例查询不需要 projectId 和 bizdate。')
fields('QueryLogsRequest', True, ['projectId','taskId','fileName','instanceId','runNumber','bizdate','projectEnv','containerId','scene','pageNumber','pageSize','maxPages'])
h('节点标识与精确匹配')
p('请求字段 nodeId 是 taskId 的兼容别名，只能填写数值型调度任务 ID。ListNodes 返回的数据开发节点 id 是 String，不可当作 taskId 使用。建议新调用方统一使用 taskId。')
p('文件名查询会分页查找并要求名称唯一且完全相等。同名节点返回 409 AMBIGUOUS_NODE，可使用 containerId、scene 缩小范围或改用 taskId。节点尚无调度任务时返回 NODE_HAS_NO_TASK。')
p('bizdate 格式为 yyyy-MM-dd，默认按 Asia/Shanghai 转为当日零点的毫秒时间戳。runNumber 只允许与 instanceId 一起传入，省略时读取云端返回的最近一次运行日志。')

page('7 日志响应字段')
h('响应中的 data')
fields('QueryLogsResponse', names=['taskId','nodeId','totalCount','pagesFetched','nextPageNumber','hasMore','complete','logs'])
h('logs 数组中的单条记录')
fields('InstanceLogResponse', names=['instanceId','taskId','taskName','runNumber','status','bizdate','startedTime','finishedTime','logContent','logSuccess','errorCode','errorMessage','requestId'])
p('complete 表示本次返回的日志是否全部读取成功，不表示已读完所有页；是否需要续读由 hasMore 判断。日志读取成功也不代表任务执行成功，任务结果应查看 status。')

page('8 日志查询示例与分页')
h('按指定实例查询')
code({'instanceId':9000001,'runNumber':1})
h('按调度任务查询')
code({'projectId':10000,'taskId':2000001,'bizdate':'2026-10-21','projectEnv':'Prod','pageNumber':1,'pageSize':20,'maxPages':1})
p('按名称查询时，将 taskId 替换为 "fileName": "python_demo"，其余筛选和分页字段保持相同。不要同时传入两种选择条件。')
h('成功响应示例')
code(ok({'taskId':2000001,'nodeId':None,'totalCount':30,'pagesFetched':1,'nextPageNumber':2,'hasMore':True,'complete':True,'logs':[{'instanceId':9000001,'taskId':2000001,'taskName':'python_demo','runNumber':1,'status':'Success','bizdate':1792512000000,'startedTime':None,'finishedTime':None,'logContent':'hello DataWorks','logSuccess':True,'errorCode':None,'errorMessage':None,'requestId':'aliyun-request-id'}]}))
p('hasMore=true 时，保留筛选条件和 pageSize，将 pageNumber 设为 nextPageNumber 继续查询。同一实例的不同 runNumber 分别保留。')
p('部分日志读取失败时 complete=false，失败记录保留 errorCode、errorMessage 和 requestId；可按 instanceId/runNumber 重查。节点查询或实例列表查询失败时，整个请求返回错误。')

page('9 错误处理与联调说明')
table(['HTTP', '常见结果码', '处理建议'], [
    ['400','INVALID_REQUEST\nINVALID_JSON\nPROJECT_ID_REQUIRED','检查必填字段、正数 ID、JSON 类型、业务日期和选择条件。未知 JSON 字段也会被拒绝。'],
    ['404','NODE_NOT_FOUND','确认文件名称及查询范围。'],
    ['409','AMBIGUOUS_NODE\nNODE_HAS_NO_TASK\nNOT_A_SUBMISSION_PACKAGE\nFILE_DEPLOYMENT_MISMATCH','检查节点是否唯一、是否已有调度任务，以及文件与提交包是否对应。'],
    ['422','NODE_SEARCH_LIMIT','名称搜索达到页数上限；缩小查询范围或直接使用 taskId。'],
    ['502','云端错误码或\nINVALID_UPSTREAM_RESPONSE\nINVALID_UPSTREAM_PAGE','记录 code 和 requestId，核对云端权限、数据及 API 版本。'],
    ['503','DATAWORKS_NOT_CONFIGURED\nREQUEST_INTERRUPTED','检查服务端凭据配置或请求中断原因。'],
    ['504','DATAWORKS_TIMEOUT\nLOG_QUERY_TIME_LIMIT','写操作先核对云端执行结果，避免盲目重试；日志查询减少每页数量。'],
    ['500','INTERNAL_ERROR','查看服务日志并排查未预期异常。']], [.5,2.65,4.05])
h('错误响应示例')
code({'success':False,'code':'INVALID_JSON','message':'JSON 格式或字段类型错误，日期格式为 yyyy-MM-dd','requestId':None,'data':None})
h('联调前确认')
p('服务端需要配置阿里云凭据、区域、工作空间及必要的 endpoint；凭据不放入接口请求体。客户端延迟初始化，缺少凭据不影响查看 Swagger，但实际业务调用可能返回 503。')
p('SDK 单次调用等待上限默认 20 秒，自动重试关闭。创建、提交或发布超时后，云端可能已经执行成功，应先核对云端状态再决定是否重试。')
p('日志聚合预算为 60 秒，在每次下游请求前检查，实际耗时可能额外包含一次进行中的 SDK 调用。在线分页不保证跨请求快照一致性。')
p('本项目面向 Java 8，使用 DataWorks SDK 9.0.11。内部状态查询调用 GetDeploymentPackage，实例和日志分别调用 ListTaskInstances、GetTaskInstanceLog；不应直接套用旧版 Blink SDK 的模型名称。')
output = ROOT / 'DataWorks接口文档.docx'
doc.save(output)
print(output)
